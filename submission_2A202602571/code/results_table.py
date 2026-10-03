"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền experiments.xlsx từ mẫu.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

FORMULA_COLS = ("step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise")
GROUP_ORDER = ("baseline", "loss", "optimizer", "hparam", "dropout", "clipping", "amp", "init", "final", "other")
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _clean(o):
    """JSON chuẩn không có NaN/inf: đổi thành None."""
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_clean(v) for v in o]
    if isinstance(o, float) and not math.isfinite(o):
        return None
    return o


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (+ notes nếu có; KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    payload = _clean({k: result[k] for k in ("cfg", "history", "summary")})
    payload["notes"] = result.get("notes", "")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file kết quả thí nghiệm (*.json có khoá cfg) trong results_dir, sắp theo exp_id."""
    out = []
    for p in sorted(Path(results_dir).glob("*.json")):
        with open(p, encoding="utf-8") as f:
            d = json.load(f)
        if isinstance(d, dict) and "cfg" in d and "summary" in d:   # bỏ qua eval_result_*.json
            out.append(d)
    return sorted(out, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một kết quả -> một dòng của bảng. Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg, s = result["cfg"], result["summary"]
    wd = cfg["weight_decay"]
    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg["description"],
        loss=cfg["loss"].upper(), optimizer=OPT_NAMES.get(cfg["optimizer"], cfg["optimizer"]),
        lr=cfg["lr"], weight_decay=wd, batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(str(h) for h in cfg["hidden"]), dropout=cfg["dropout"],
        clip_norm="none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        figure_file=f"figures/{cfg['exp_id']}.png",
    )
    for k in ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
              "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB"):
        row[k] = s.get(k)                      # None -> ô trống (lần chạy phân kỳ / không đo được)
    row["diverged"] = "Có" if s.get("diverged") else "Không"
    if eval_scores:
        row["eval_acc"], row["eval_macro_f1"] = eval_scores["accuracy"], eval_scores["macro_f1"]
    parts = [result.get("notes", ""), notes]
    if cfg["optimizer"] == "sgd_momentum":
        parts.append(f"momentum={cfg['momentum']}")
    if cfg["optimizer"] in ("adam", "adamw"):
        parts.append("betas=(0.9,0.999), eps=1e-8")
    if cfg.get("scheduler"):
        parts.append(f"scheduler={cfg['scheduler']} (1 epoch khởi động)")
    if s.get("clip_fraction") is not None:
        parts.append(f"clip kích hoạt ở {100 * s['clip_fraction']:.1f}% số bước")
    if s.get("diverged"):
        parts.append(f"phân kỳ: {s.get('diverged_at')}")
    if s.get("peak_mem_MB") is None:
        parts.append("peak_mem không đo (không có CUDA)")
    row["notes"] = "; ".join(p for p in parts if p)
    return row


def sort_rows(rows: list[dict]) -> list[dict]:
    """base-s1 lên đầu (dòng 2 của mẫu là baseline), sau đó theo thứ tự nhóm rồi exp_id."""
    def key(r):
        g = GROUP_ORDER.index(r["group"]) if r["group"] in GROUP_ORDER else len(GROUP_ORDER)
        return (g, r["exp_id"])
    return sorted(rows, key=key)


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền sheet "Experiments" từ dòng 2; tuỳ chọn điền exp_id baseline vào sheet "Seeds" (cột A)
    và nhận xét theo nhóm vào sheet "Summary" (cột H). Không đụng tới ô công thức, tên sheet, tên cột.

    openpyxl không tính công thức: mở file bằng Excel/LibreOffice rồi lưu để giá trị được tính lại.
    """
    import openpyxl
    wb = openpyxl.load_workbook(template_path)          # KHÔNG data_only=True (sẽ mất công thức)
    ws = wb["Experiments"]
    col = {c.value: c.column for c in ws[1] if c.value}
    capacity = ws.max_row - 1
    if len(rows) > capacity:
        raise ValueError(f"mẫu chỉ có {capacity} dòng có công thức, cần {len(rows)}")
    for i, row in enumerate(rows):
        for name, c in col.items():
            if name in FORMULA_COLS:
                continue
            ws.cell(row=2 + i, column=c).value = row.get(name)
    if seed_ids is not None:
        wsd = wb["Seeds"]
        if len(seed_ids) > 5:
            raise ValueError("sheet Seeds chỉ có 5 dòng (A2:A6)")
        for i in range(5):
            wsd.cell(row=2 + i, column=1).value = seed_ids[i] if i < len(seed_ids) else None
    if summary_notes:
        wss = wb["Summary"]
        for r in range(2, wss.max_row + 1):
            g = wss.cell(row=r, column=1).value
            if g in summary_notes:
                wss.cell(row=r, column=8).value = summary_notes[g]
    wb.save(out_path)
